"""EXTENSÃO — relatório da qualidade da entrada.

    python repro/51_quality_report.py signals   # sinais objetivos, sem juiz (segundos)
    python repro/51_quality_report.py sample    # planilha com 100 issues para VOCÊ rotular
    python repro/51_quality_report.py report    # resumo do juiz, concordância e erro × qualidade

signals  Medidas que não dependem de nenhum modelo: descrição vazia, número de palavras,
         formato "As a ... I want", presença de "so that" e de critérios de aceitação.
         Saída: results/quality/sinais_objetivos.csv (uma linha por issue).

sample   Sorteia 100 issues de avaliação (50 Agile + 50 NEO) e gera
         results/quality/rotulagem_humana.xlsx, com colunas em branco para as mesmas
         perguntas do juiz. O artigo do Jev recomenda ~100 rótulos locais para decidir se
         dá para confiar no juiz e escolher o limiar de confiança. NÃO rode de novo depois de
         começar a rotular (sobrescreve; use --force se for de propósito).

report   (1) distribuição das respostas do juiz;
         (2) concordância juiz × seus rótulos (acurácia e kappa de Cohen), e a curva
             "aceitar quando confiante": acurácia e cobertura para cada limiar de confiança;
         (3) erro dos modelos por nível de qualidade (sinais objetivos e juiz), usando as
             predições já salvas em results/predictions/*/within/.
         Saídas em results/quality/.
"""
import argparse
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from see import config  # noqa: E402
from see.io import load_merged  # noqa: E402
from see.quality import add_quality_scores, summarize  # noqa: E402

OUT = config.RESULTS / "quality"
LABEL_FILE = OUT / "rotulagem_humana.xlsx"
QUESTIONS = ["issue_type", "goal_clarity", "estimability", "scope_size", "has_acceptance_criteria", "is_ambiguous"]
SCORE_Q = ["goal_clarity", "estimability", "scope_size"]
NOUL_Q = ["has_acceptance_criteria", "is_ambiguous"]

US_FORMAT = re.compile(r"\bas an?\s+.{1,60}?\bi\s+(?:want|would like|need|can|should)\b", re.I | re.S)
SO_THAT = re.compile(r"\bso that\b", re.I)
ACCEPTANCE = re.compile(r"acceptance criteria|given .{1,80} when .{1,80} then|\[ \]|\[x\]|definition of done",
                        re.I | re.S)


# ---------------------------------------------------------------- signals
def compute_signals(df: pd.DataFrame) -> pd.DataFrame:
    desc = df.description.fillna("").astype(str)
    words = df.text_light.fillna("").str.split().str.len()
    return pd.DataFrame({
        "issue_id": df.issue_id, "dataset": df.dataset, "project": df.project, "sp": df.sp,
        "within_eval": df.within_eval,
        "empty_description": (desc.str.strip() == "").astype(int),
        "n_words": words,
        "under_15_words": (words < 15).astype(int),
        "user_story_format": df.text_light.fillna("").str.contains(US_FORMAT).astype(int),
        "has_so_that": df.text_light.fillna("").str.contains(SO_THAT).astype(int),
        "acceptance_criteria_text": desc.str.contains(ACCEPTANCE).astype(int),
    })


def cmd_signals(_args):
    OUT.mkdir(parents=True, exist_ok=True)
    sig = pd.concat([compute_signals(load_merged(d)) for d in config.DATASETS], ignore_index=True)
    sig.to_csv(OUT / "sinais_objetivos.csv", index=False)
    cols = ["empty_description", "under_15_words", "user_story_format", "has_so_that", "acceptance_criteria_text"]
    summary = []
    for (ds, scope), g in [((d, "todas"), sig[sig.dataset == d]) for d in config.DATASETS] + \
                          [((d, "avaliação"), sig[(sig.dataset == d) & (sig.within_eval == 1)]) for d in config.DATASETS]:
        row = {"dataset": ds, "issues": scope, "n": len(g), "mediana_palavras": g.n_words.median()}
        row.update({c: round(100 * g[c].mean(), 1) for c in cols})
        summary.append(row)
    summary = pd.DataFrame(summary)
    summary.to_csv(OUT / "sinais_objetivos_resumo.csv", index=False)
    print("% das issues com cada sinal:")
    print(summary.to_string(index=False))
    print(f"\nSalvo em {OUT.relative_to(config.ROOT)}/")


# ---------------------------------------------------------------- sample
def cmd_sample(args):
    if LABEL_FILE.exists() and not args.force:
        sys.exit(f"{LABEL_FILE.relative_to(config.ROOT)} já existe (pode ter rótulos seus). Use --force para sobrescrever.")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    rng = np.random.default_rng(config.SEED)
    parts = []
    for d in config.DATASETS:
        ev = load_merged(d)
        ev = ev[ev.within_eval == 1]
        parts.append(ev.iloc[rng.choice(len(ev), args.n // 2, replace=False)])
    sample = pd.concat(parts)
    wb = Workbook()
    ws = wb.active
    ws.title = "rotular"
    head = ["issue_id", "dataset", "title", "description",
            "issue_type", "goal_clarity", "estimability", "scope_size", "has_acceptance_criteria", "is_ambiguous",
            "observacao"]
    ws.append(head)
    for r in sample.itertuples():
        ws.append([r.issue_id, r.dataset, r.title, str(r.description)[:3000]] + [None] * 7)
    bold = Font(name="Arial", bold=True)
    fill = PatternFill("solid", fgColor="FFFF00")
    for c in ws[1]:
        c.font = bold
    for col in "EFGHIJ":
        for cell in ws[col][1:]:
            cell.fill = fill
    for col, w in zip("ABCDEFGHIJK", [22, 8, 40, 90, 16, 13, 13, 12, 14, 12, 30]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row[2:4]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    n = len(sample) + 1
    validations = {
        "E": '"user_story,bug,technical_task,improvement,other"',
        "F": '"0,1,2"', "G": '"0,1,2"', "H": '"0,1,2"', "I": '"0,1"', "J": '"0,1"',
    }
    for col, formula in validations.items():
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    ws.freeze_panes = "E2"

    guide = wb.create_sheet("instrucoes")
    lines = [
        ("Como rotular", "Leia título e descrição e preencha as colunas amarelas (há listas suspensas). "
                         "Responda como se fosse estimar a issue, sem olhar os story points."),
        ("issue_type", "user_story = funcionalidade nova do ponto de vista do usuário; bug = defeito; "
                       "technical_task = trabalho interno (refatorar, infraestrutura, testes, dependências); "
                       "improvement = melhoria do que já existe; other = pergunta, discussão, documentação."),
        ("goal_clarity", "0 = não dá para saber o que mudar; 1 = dá para entender fazendo suposições; 2 = a mudança está explícita."),
        ("estimability", "0 = falta informação essencial para estimar; 1 = parcial; 2 = dá para estimar com segurança razoável."),
        ("scope_size", "0 = pequena e bem delimitada; 1 = várias mudanças relacionadas; 2 = grande, aberta ou do tamanho de um épico."),
        ("has_acceptance_criteria", "1 = diz como verificar que está pronto (critérios de aceitação, passos esperados); 0 = não."),
        ("is_ambiguous", "1 = termos vagos que pessoas diferentes interpretariam de jeitos diferentes; 0 = não."),
        ("Depois", "Salve este arquivo no mesmo lugar e rode: python repro\\51_quality_report.py report"),
    ]
    for a, b in lines:
        guide.append([a, b])
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 120
    for row in guide.iter_rows():
        row[0].font = bold
        row[1].alignment = Alignment(wrap_text=True)
    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(LABEL_FILE)
    print(f"{len(sample)} issues para rotular em {LABEL_FILE.relative_to(config.ROOT)}")


# ---------------------------------------------------------------- report
def load_judge():
    frames = []
    for f in sorted(OUT.glob("*__*.csv")):
        dataset, engine = f.stem.split("__", 1)
        d = pd.read_csv(f, dtype={"issue_id": str})
        d["dataset"], d["engine"] = dataset, engine
        frames.append(add_quality_scores(d))
    return pd.concat(frames, ignore_index=True) if frames else None


def load_labels():
    if not LABEL_FILE.exists():
        return None
    lab = pd.read_excel(LABEL_FILE, sheet_name="rotular", dtype={"issue_id": str})
    lab = lab.dropna(subset=QUESTIONS, how="all")
    return lab if len(lab) else None


def agreement(judge: pd.DataFrame, labels: pd.DataFrame) -> pd.DataFrame:
    from sklearn.metrics import cohen_kappa_score
    rows, curves = [], []
    for engine, j in judge.groupby("engine"):
        m = labels.merge(j, on="issue_id", suffixes=("_h", "_j"))
        for q in QUESTIONS:
            h_col, j_col = f"{q}_h", f"{q}_j"
            if h_col not in m or j_col not in m:
                continue
            x = m[[h_col, j_col] + ([f"{q}_conf"] if f"{q}_conf" in m else [])].dropna(subset=[h_col, j_col])
            if len(x) < 5:
                continue
            h = x[h_col]
            if q in SCORE_Q:
                jv = x[j_col].round().clip(0, 2).astype(int)
                h = h.astype(int)
                kappa = cohen_kappa_score(h, jv, weights="quadratic")
            elif q in NOUL_Q:
                jv = (x[j_col] >= 0.5).astype(int)
                h = h.astype(int)
                kappa = cohen_kappa_score(h, jv)
            else:
                jv = x[j_col].astype(str)
                h = h.astype(str).str.strip()
                kappa = cohen_kappa_score(h, jv)
            hit = (h.values == jv.values)
            rows.append({"engine": engine, "pergunta": q, "n": len(x), "acuracia": hit.mean(), "kappa": kappa,
                         "kappa_tipo": "ponderado (quadrático)" if q in SCORE_Q else "Cohen"})
            conf_col = f"{q}_conf"
            if conf_col in x:
                for tau in [0.0, 0.5, 0.6, 0.7, 0.8, 0.9]:
                    keep = (x[conf_col] >= tau).values
                    if keep.sum():
                        curves.append({"engine": engine, "pergunta": q, "limiar": tau,
                                       "cobertura": keep.mean(), "acuracia_aceitas": hit[keep].mean()})
    return pd.DataFrame(rows), pd.DataFrame(curves)


def error_by_quality(judge, signals) -> pd.DataFrame:
    """MAE de cada modelo (within) por nível de qualidade."""
    out = []
    for dataset in config.DATASETS:
        pdir = config.PREDICTIONS / dataset / "within"
        if not pdir.exists():
            continue
        q = signals[signals.dataset == dataset][["issue_id", "empty_description", "under_15_words"]].copy()
        if judge is not None:
            jd = judge[judge.dataset == dataset]
            for engine, j in jd.groupby("engine"):
                cols = (["issue_id"] + [c for c in SCORE_Q if c in j]
                        + [c for c in ["issue_type", "faixa"] if c in j])
                jj = j[cols].copy()
                for c in SCORE_Q:
                    if c in jj:
                        jj[c] = jj[c].round().clip(0, 2).astype("Int64")
                jj = jj.rename(columns={c: f"{engine}:{c}" for c in cols if c != "issue_id"})
                q = q.merge(jj, on="issue_id", how="left")
        dims = [c for c in q.columns if c != "issue_id"]
        for f in sorted(pdir.glob("*.csv")):
            p = pd.read_csv(f, dtype={"issue_id": str})
            p = p.merge(q, on="issue_id")
            p = p[p.issue_id.isin(signals.loc[(signals.dataset == dataset) & (signals.within_eval == 1), "issue_id"])]
            p["abs_err"] = (p.y_true - p.y_pred).abs()
            for dim in dims:
                for level, g in p.dropna(subset=[dim]).groupby(dim):
                    out.append({"dataset": dataset, "modelo": f.stem, "dimensao": dim, "nivel": level,
                                "n": len(g), "MAE": g.abs_err.mean(), "SP_medio": g.y_true.mean()})
    return pd.DataFrame(out)


def cmd_report(_args):
    OUT.mkdir(parents=True, exist_ok=True)
    sig_path = OUT / "sinais_objetivos.csv"
    if not sig_path.exists():
        cmd_signals(None)
    signals = pd.read_csv(sig_path, dtype={"issue_id": str})
    judge = load_judge()

    if judge is not None:
        print("=== (0) Nota geral de qualidade (0–100) ===")
        resumo = pd.DataFrame([{"dataset": ds, "engine": eng, **summarize(g)}
                               for (ds, eng), g in judge.groupby(["dataset", "engine"])])
        resumo.to_csv(OUT / "nota_geral_resumo.csv", index=False)
        print(resumo.to_string(index=False))
        print("indice = média de clareza, dá-para-estimar, critérios de aceitação e (1 − ambiguidade); "
              "nota_juiz = avaliação holística do juiz. Faixas: ruim < 40 ≤ regular < 70 ≤ boa.")
        if "nota_geral_juiz" in judge:
            print("spearman_indice_x_nota_juiz perto de 1 = as duas notas concordam; perto de 0 = não.")
        print()
        print("=== (1) Respostas do juiz ===")
        for (ds, eng), g in judge.groupby(["dataset", "engine"]):
            print(f"\n[{ds} / {eng}] {len(g)} issues, latência mediana {g.latency_s.median():.3f}s")
            if "issue_type" in g:
                print("  tipo:", (g.issue_type.value_counts(normalize=True) * 100).round(1).to_dict())
            for c in SCORE_Q:
                if c in g:
                    dist = (g[c].round().value_counts(normalize=True).sort_index() * 100).round(1).to_dict()
                    print(f"  {c}: média {g[c].mean():.2f} (0–2) | níveis %: {dist}")
            for c in NOUL_Q:
                if c in g:
                    print(f"  {c}: {100 * (g[c] >= 0.5).mean():.1f}% sim")
    else:
        print("Nenhum resultado do juiz ainda (rode 50_quality_judge.py). Seguindo só com os sinais objetivos.")

    labels = load_labels()
    if labels is not None and judge is not None:
        agree, curves = agreement(judge, labels)
        agree.to_csv(OUT / "concordancia_juiz_humano.csv", index=False)
        curves.to_csv(OUT / "limiar_confianca.csv", index=False)
        print(f"\n=== (2) Concordância com seus {len(labels)} rótulos ===")
        print(agree.round(3).to_string(index=False))
        print("\nKappa: < 0,2 fraco; 0,2–0,4 razoável; 0,4–0,6 moderado; 0,6–0,8 bom; > 0,8 muito bom.")
        if len(curves):
            print("\nAceitar só quando confiante (cobertura × acurácia):")
            print(curves.round(3).to_string(index=False))
    elif labels is None:
        print("\n(2) Sem rótulos humanos ainda: rode `sample`, preencha a planilha e rode `report` de novo.")

    err = error_by_quality(judge, signals)
    if len(err):
        err.to_csv(OUT / "erro_por_qualidade.csv", index=False)
        print("\n=== (3) MAE por nível de qualidade (issues de avaliação) ===")
        show = err[err.modelo.isin(["ProjectMedian", "XGBoost", "SVR", "RandomForest"]) |
                   err.modelo.str.contains("__ZS1|BERT|GPT2SP", regex=True)]
        piv = show.pivot_table(index=["dataset", "dimensao", "nivel"], columns="modelo", values="MAE").round(2)
        n = show.groupby(["dataset", "dimensao", "nivel"]).n.max()
        piv.insert(0, "n", n)
        print(piv.to_string())
        print(f"\nTabela completa: {(OUT / 'erro_por_qualidade.csv').relative_to(config.ROOT)}")


# ---------------------------------------------------------------- inspect
def describe_answers(r) -> list:
    lines = []
    if "indice_qualidade" in r and pd.notna(r.get("indice_qualidade")):
        extra = f" | nota geral do juiz: {r['nota_geral_juiz']:.0f}/100" if pd.notna(r.get("nota_geral_juiz", np.nan)) else ""
        lines.append(f"ÍNDICE DE QUALIDADE: {r['indice_qualidade']:.0f}/100 ({r.get('faixa')}){extra}")
    if "issue_type" in r and pd.notna(r.get("issue_type")):
        probs = {k[len("issue_type_p_"):]: r[k] for k in r.index if k.startswith("issue_type_p_") and pd.notna(r[k])}
        top = ", ".join(f"{k} {v:.0%}" for k, v in sorted(probs.items(), key=lambda kv: -kv[1])[:3])
        lines.append(f"tipo: {r.issue_type} (confiança {r.issue_type_conf:.2f}; {top})")
    names = {"goal_clarity": "clareza", "estimability": "dá para estimar", "scope_size": "escopo (0 pequeno, 2 grande)"}
    for q, label in names.items():
        if q in r and pd.notna(r.get(q)):
            lines.append(f"{label}: {r[q]:.2f} de 2 (confiança {r.get(q + '_conf', float('nan')):.2f})")
    for q, label in {"has_acceptance_criteria": "tem critérios de aceitação", "is_ambiguous": "é ambígua"}.items():
        if q in r and pd.notna(r.get(q)):
            lines.append(f"{label}: {r[q]:.0%} -> {'SIM' if r[q] >= 0.5 else 'não'}")
    return lines


def cmd_inspect(args):
    path = OUT / f"{args.dataset}__{args.engine}.csv"
    if not path.exists():
        sys.exit(f"{path.relative_to(config.ROOT)} não existe. Rode antes o 50_quality_judge.py.")
    j = add_quality_scores(pd.read_csv(path, dtype={"issue_id": str}))
    data = load_merged(args.dataset).set_index("issue_id")
    j = j.head(args.n) if args.n else j
    rows = []
    for _, r in j.iterrows():
        d = data.loc[r.issue_id]
        desc = str(d.description or "").replace("\n", " ")
        print("=" * 100)
        print(f"{r.issue_id} | projeto {d.project} | story points {d.sp:g}")
        print(f"TÍTULO: {d.title}")
        print(f"DESCRIÇÃO: {desc[:600]}{' ...' if len(desc) > 600 else ''}")
        print("-- juiz --")
        answers = describe_answers(r)
        for line in answers:
            print("  " + line)
        rows.append({"issue_id": r.issue_id, "story_points": d.sp, "titulo": d.title, "descricao": desc[:3000],
                     "resposta_do_juiz": "\n".join(answers), "concorda? (s/n)": None, "comentario": None})
    if args.xlsx:
        from openpyxl.styles import Alignment, Font, PatternFill
        out = OUT / f"conferencia_{args.dataset}_{args.engine}.xlsx"
        if out.exists() and not args.force:
            sys.exit(f"{out.relative_to(config.ROOT)} já existe (pode ter marcações suas). Use --force para sobrescrever.")
        df = pd.DataFrame(rows)
        with pd.ExcelWriter(out, engine="openpyxl") as w:
            df.to_excel(w, index=False, sheet_name="conferir")
            ws = w.sheets["conferir"]
            for col, width in zip("ABCDEFG", [16, 8, 40, 80, 55, 12, 30]):
                ws.column_dimensions[col].width = width
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                row[5].fill = PatternFill("solid", fgColor="FFFF00")
            for c in ws[1]:
                c.font = Font(name="Arial", bold=True)
        print(f"\nPlanilha para conferir: {out.relative_to(config.ROOT)}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("comando", choices=["signals", "sample", "report", "inspect"])
    ap.add_argument("--dataset", default="agile", choices=config.ALL_DATASETS, help="(inspect)")
    ap.add_argument("--engine", default="jev", help="(inspect) jev ou laya")
    ap.add_argument("--xlsx", action="store_true", help="(inspect) gerar planilha de conferência")
    ap.add_argument("--n", type=int, default=None,
                    help="(sample) total de issues, metade de cada dataset (padrão 100); (inspect) quantas mostrar (padrão todas)")
    ap.add_argument("--force", action="store_true", help="(sample) sobrescrever a planilha de rótulos")
    args = ap.parse_args()
    if args.comando == "sample" and args.n is None:
        args.n = 100
    {"signals": cmd_signals, "sample": cmd_sample, "report": cmd_report, "inspect": cmd_inspect}[args.comando](args)


if __name__ == "__main__":
    main()
