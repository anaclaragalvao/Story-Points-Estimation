"""Gera results/tables/comparacao_artigo_vs_reproducao.xlsx (artigo × reprodução).

    python repro/90_evaluate.py                 # atualiza metrics.csv
    python repro/92_build_comparison_xlsx.py    # gera a planilha
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pathlib import Path
_TABLES = Path(__file__).resolve().parent / "results" / "tables"
METRICS = sys.argv[1] if len(sys.argv) > 1 else str(_TABLES / "metrics.csv")
OUT = sys.argv[2] if len(sys.argv) > 2 else str(_TABLES / "comparacao_artigo_vs_reproducao.xlsx")
m = pd.read_csv(METRICS)

FONT = "Arial"
F = lambda **k: Font(name=FONT, size=10, **k)  # noqa: E731
BLUE = F(color="0000FF")
BOLD = F(bold=True)
TITLE = Font(name=FONT, size=13, bold=True)
WHITE_B = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
ART_FILL = PatternFill("solid", fgColor="FCE4D6")
OURS_FILL = PatternFill("solid", fgColor="E2EFDA")
CODE_FILL = PatternFill("solid", fgColor="FFF2CC")
FAM_FILL = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)

SOURCE_PAPER = ("Alaswad, Poovammal & Aljaddouh, 'From COCOMO to GPT', IEEE Access 14 (2026), "
                "DOI 10.1109/ACCESS.2026.3671204")

# ---------------- valores do artigo (Tabelas 14 e 15) ----------------
PAPER = {
    "agile": {
        "SVR": (4.257, 2.321, 10.066, 0.230), "Random Forest": (5.112, 3.610, 10.134, 0.191),
        "XGBoost": (4.971, 3.177, 10.050, 0.205), "KNN": (5.268, 3.403, 10.380, 0.193),
        "Linear Regression": (5.934, 3.837, 10.721, 0.153),
        "BERT": (4.112, 2.125, 8.203, 0.176), "RoBERTa": (3.781, 1.982, 7.544, 0.232),
        "GPT-2": (3.478, 1.719, 7.538, 0.279), "GPT2SP": (2.860, 1.402, 6.921, 0.327),
        "GPT-4-Turbo": (7.215, 4.386, 12.984, 0.082), "DeepSeek-V3": (8.063, 5.022, 14.673, 0.061),
        "LLaMA-3": (9.148, 5.815, 16.839, 0.044), "Gemini-1.5": (10.393, 6.772, 18.974, 0.029),
    },
    "neo": {
        "SVR": (5.684, 3.173, 11.482, 0.182), "Random Forest": (6.715, 4.383, 12.913, 0.138),
        "XGBoost": (6.032, 3.598, 11.937, 0.167), "KNN": (6.387, 4.032, 12.347, 0.153),
        "Linear Regression": (7.184, 4.714, 13.605, 0.119),
        "BERT": (4.982, 2.817, 9.413, 0.152), "RoBERTa": (4.637, 2.593, 8.973, 0.176),
        "GPT-2": (4.292, 2.342, 8.721, 0.202), "GPT2SP": (3.705, 2.041, 7.986, 0.231),
        "GPT-4-Turbo": (8.913, 5.732, 15.943, 0.072), "DeepSeek-V3": (11.917, 7.032, 18.946, 0.042),
        "LLaMA-3": (9.842, 6.383, 17.318, 0.056), "Gemini-1.5": (12.184, 7.915, 21.387, 0.028),
    },
}
FAMILY = {"SVR": "ML", "Random Forest": "ML", "XGBoost": "ML", "KNN": "ML", "Linear Regression": "ML",
          "BERT": "Transformer", "RoBERTa": "Transformer", "GPT-2": "Transformer", "GPT2SP": "Transformer",
          "GPT-4-Turbo": "LLM", "DeepSeek-V3": "LLM", "LLaMA-3": "LLM", "Gemini-1.5": "LLM"}
# nome no artigo -> nome do arquivo de predições na reprodução
OURS_NAME = {"SVR": "SVR", "Random Forest": "RandomForest", "XGBoost": "XGBoost", "KNN": "KNN",
             "Linear Regression": "LinearRegression"}


def ours(dataset, protocol, model):
    r = m[(m.dataset == dataset) & (m.protocol == protocol) & (m.model == model)]
    if r.empty:
        return None
    r = r.iloc[0]
    return (r.MAE, r.MdAE, r.RMSE, r.PRED25, int(r.n))


def status(dataset, name, art, code):
    fam = FAMILY[name]
    if fam == "LLM":
        return "Pendente: rodar 30_run_llm.py (os modelos exatos do artigo não são gratuitos; usaremos equivalentes abertos)"
    if fam == "Transformer":
        if name == "GPT2SP":
            return "Pendente: Colab (20_run_transformers.py). O GPT2SP não existe no código original."
        return "Pendente: Colab (20_run_transformers.py e, no Agile, 21_run_transformers_original_code.py)"
    if name == "Linear Regression" and art and art[0] > 20:
        return "Sem SVD a regressão degenera (mais features que exemplos); o artigo só é plausível com a SVD do código"
    if dataset == "agile" and code:
        return "Código original reproduz o artigo; pelo protocolo descrito no texto, o erro fica menor"
    if dataset == "agile" and name == "SVR":
        return "Código original: SVR pendente (> 1 h de CPU)"
    if dataset == "neo" and art is None:
        return "Pendente: python repro\\10_run_ml.py --dataset neo --models SVR (~30 min)"
    if dataset == "neo":
        return "Sem código do artigo para o NEO; pelo protocolo do texto, o erro fica menor"
    return ""


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_range(ws, row, c1, c2, font=None, fill=None, align=None, border=True, fmt=None):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        if font: cell.font = font
        if fill: cell.fill = fill
        if align: cell.alignment = align
        if border: cell.border = BOX
        if fmt: cell.number_format = fmt


def comparison_sheet(wb, dataset, title, with_code):
    ws = wb.create_sheet(title)
    label = "Agile User Story Point Dataset" if dataset == "agile" else "NEO Dataset"
    ws["A1"] = f"{label} — story points: artigo × reprodução"
    ws["A1"].font = TITLE
    note = ("Artigo = Tabela 14" if dataset == "agile" else "Artigo = Tabela 15") + \
           " (protocolo within-project). Valores em azul são dados copiados; colunas de diferença são fórmulas."
    ws["A2"] = note
    ws["A2"].font = F(italic=True)

    groups = [("Artigo (publicado)", ART_FILL, 4),
              ("Reprodução — protocolo descrito no artigo", OURS_FILL, 5)]
    if with_code:
        groups.append(("Reprodução — código original executado", CODE_FILL, 5))
    groups.append(("Diferença MAE (nós − artigo)", SUB_FILL, 4 if with_code else 2))
    groups.append(("Situação", SUB_FILL, 1))

    r = 4
    ws.cell(row=r, column=1, value="Família"); ws.cell(row=r, column=2, value="Modelo")
    style_range(ws, r, 1, 2, font=WHITE_B, fill=HDR_FILL, align=CENTER)
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
    ws.merge_cells(start_row=r, start_column=2, end_row=r + 1, end_column=2)
    col = 3
    sub = []
    for name, fill, width in groups:
        ws.cell(row=r, column=col, value=name)
        style_range(ws, r, col, col + width - 1, font=WHITE_B, fill=HDR_FILL, align=CENTER)
        if width > 1:
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + width - 1)
        sub.append((col, name, fill, width))
        col += width
    last_col = col - 1

    r2 = r + 1
    heads = []
    for start, name, fill, width in sub:
        if name.startswith("Artigo"):
            labels = ["MAE", "MdAE", "RMSE", "PRED(25)"]
        elif name.startswith("Reprodução"):
            labels = ["MAE", "MdAE", "RMSE", "PRED(25)", "n issues"]
        elif name.startswith("Diferença"):
            labels = ["Protocolo do artigo", "%"] + (["Código original", "%"] if with_code else [])
        else:
            labels = [""]
        for i, lab in enumerate(labels):
            ws.cell(row=r2, column=start + i, value=lab)
        style_range(ws, r2, start, start + width - 1, font=BOLD, fill=fill, align=CENTER)
        heads.append((start, name))
    # posições das colunas
    c_art = 3
    c_ours = c_art + 4
    c_code = c_ours + 5 if with_code else None
    c_diff = (c_code + 5) if with_code else (c_ours + 5)
    c_status = last_col

    row = r2 + 1
    first_data = row
    prev_fam = None
    for name, art in PAPER[dataset].items():
        fam = FAMILY[name]
        ws.cell(row=row, column=1, value=fam if fam != prev_fam else None)
        prev_fam = fam
        ws.cell(row=row, column=2, value=name)
        style_range(ws, row, 1, last_col, font=F(), align=Alignment(vertical="center"))
        ws.cell(row=row, column=1).font = BOLD
        for i, v in enumerate(art):
            c = ws.cell(row=row, column=c_art + i, value=v)
            c.font = BLUE
            c.number_format = "0.0%" if i == 3 else "0.000"
        o = ours(dataset, "within", OURS_NAME.get(name, "∅"))
        if o:
            for i, v in enumerate(o):
                c = ws.cell(row=row, column=c_ours + i, value=float(v) if i < 4 else v)
                c.font = BLUE
                c.number_format = "0.0%" if i == 3 else ("0" if i == 4 else "0.000")
        code = ours(dataset, "original_code", OURS_NAME.get(name, "∅")) if with_code else None
        if code:
            for i, v in enumerate(code):
                c = ws.cell(row=row, column=c_code + i, value=float(v) if i < 4 else v)
                c.font = BLUE
                c.number_format = "0.0%" if i == 3 else ("0" if i == 4 else "0.000")
        A = f"{get_column_letter(c_art)}{row}"
        Om = f"{get_column_letter(c_ours)}{row}"
        ws.cell(row=row, column=c_diff, value=f'=IF({Om}="","",{Om}-{A})').number_format = "+0.000;-0.000;0.000"
        ws.cell(row=row, column=c_diff + 1, value=f'=IF({Om}="","",({Om}-{A})/{A})').number_format = "+0.0%;-0.0%;0.0%"
        if with_code:
            Cm = f"{get_column_letter(c_code)}{row}"
            ws.cell(row=row, column=c_diff + 2, value=f'=IF({Cm}="","",{Cm}-{A})').number_format = "+0.000;-0.000;0.000"
            ws.cell(row=row, column=c_diff + 3, value=f'=IF({Cm}="","",({Cm}-{A})/{A})').number_format = "+0.0%;-0.0%;0.0%"
        st = ws.cell(row=row, column=c_status, value=status(dataset, name, o, code))
        st.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 14 * (len(st.value) // 60 + 1)
        if fam != "ML":
            style_range(ws, row, 3, c_status - 1, fill=FAM_FILL, border=True)
        row += 1
    last_data = row - 1

    # Baselines (não reportados no artigo)
    row += 1
    ws.cell(row=row, column=1, value="Baselines ingênuos (o artigo NÃO reporta) — mesmo protocolo e mesmas issues")
    ws.cell(row=row, column=1).font = BOLD
    row += 1
    base_names = {"ProjectMedian": "Mediana do projeto", "ProjectMean": "Média do projeto",
                  "Median": "Mediana global", "Mean": "Média global"}
    for key, label in base_names.items():
        o = ours(dataset, "within", key)
        ws.cell(row=row, column=1, value="Baseline" if key == "ProjectMedian" else None).font = BOLD
        ws.cell(row=row, column=2, value=label)
        style_range(ws, row, 1, c_ours + 4, font=F())
        ws.cell(row=row, column=1).font = BOLD
        for i in range(4):
            ws.cell(row=row, column=c_art + i, value="—").alignment = CENTER
        if o:
            for i, v in enumerate(o):
                c = ws.cell(row=row, column=c_ours + i, value=float(v) if i < 4 else v)
                c.font = BLUE
                c.number_format = "0.0%" if i == 3 else ("0" if i == 4 else "0.000")
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Fontes").font = BOLD
    ws.cell(row=row + 1, column=1, value=f"Artigo: {SOURCE_PAPER}, "
            + ("Tabela 14 (Apêndice C)." if dataset == "agile" else "Tabela 15 (Apêndice C)."))
    ws.cell(row=row + 2, column=1, value="Reprodução: repro/results/tables/metrics.csv, gerado por 90_evaluate.py "
            "a partir das predições em repro/results/predictions/. MAE, MdAE e RMSE em story points; "
            "PRED(25) = % de estimativas com erro relativo ≤ 25%.")
    ws.cell(row=row + 3, column=1, value="Protocolo do artigo: split 80/10/10 dentro de cada projeto, avaliação em 50% do teste "
            + ("(1.166 issues)." if dataset == "agile" else "(1.022 issues)."))
    if with_code:
        ws.cell(row=row + 4, column=1, value="Código original: split aleatório global 70/30 (6.994 issues de teste), TF-IDF de "
                "93 mil termos e SVD de 3.000 componentes ajustados no dataset inteiro, executado em float32.")
    for rr in range(row + 1, row + 5):
        ws.cell(row=rr, column=1).font = F(italic=True)

    ws.freeze_panes = ws.cell(row=first_data, column=3)
    widths = [12, 18] + [9] * (c_status - 3) + [62]
    set_widths(ws, widths)
    ws.row_dimensions[r].height = 30
    return ws


def generalization_sheet(wb):
    ws = wb.create_sheet("Generalizacao")
    ws["A1"] = "Generalização para projetos novos: artigo × reprodução"
    ws["A1"].font = TITLE
    ws["A2"] = ("Artigo = Tabela 7: variação média % do MAE por família, em relação ao within. "
                "Reprodução: mesma conta, por modelo e média da família (fórmulas).")
    ws["A2"].font = F(italic=True)

    hdr = ["Dataset", "Protocolo", "Família", "Modelo", "MAE within (nós)", "MAE no protocolo (nós)",
           "Δ MAE % (nós)", "Δ MAE % artigo (família)"]
    r = 4
    for i, h in enumerate(hdr, start=1):
        ws.cell(row=r, column=i, value=h)
    style_range(ws, r, 1, len(hdr), font=WHITE_B, fill=HDR_FILL, align=CENTER)
    paper_t7 = {("agile", "cross", "ML"): 0.081, ("agile", "lopo", "ML"): 0.130,
                ("neo", "cross", "ML"): 0.135, ("neo", "lopo", "ML"): 0.210}
    row = r + 1
    first = row
    for name in ["KNN", "RandomForest", "XGBoost"]:
        w = ours("neo", "within", name)
        c = ours("neo", "cross", name)
        if not (w and c):
            continue
        vals = ["NEO", "cross", "ML", name, float(w[0]), float(c[0])]
        for i, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            if i >= 5:
                cell.font = BLUE
                cell.number_format = "0.000"
        ws.cell(row=row, column=7, value=f"=(F{row}-E{row})/E{row}").number_format = "+0.0%;-0.0%;0.0%"
        ws.cell(row=row, column=8, value=paper_t7[("neo", "cross", "ML")]).number_format = "+0.0%;-0.0%;0.0%"
        ws.cell(row=row, column=8).font = BLUE
        style_range(ws, row, 1, len(hdr), font=None)
        row += 1
    last = row - 1
    ws.cell(row=row, column=4, value="Média ML (NEO cross)").font = BOLD
    ws.cell(row=row, column=7, value=f"=AVERAGE(G{first}:G{last})").number_format = "+0.0%;-0.0%;0.0%"
    ws.cell(row=row, column=7).font = BOLD
    ws.cell(row=row, column=8, value=paper_t7[("neo", "cross", "ML")]).number_format = "+0.0%;-0.0%;0.0%"
    ws.cell(row=row, column=8).font = BLUE
    style_range(ws, row, 1, len(hdr), fill=SUB_FILL)
    ws.cell(row=row, column=4).font = BOLD

    row += 2
    ws.cell(row=row, column=1, value="Tabela 7 do artigo completa (Δ % médio por família)").font = BOLD
    row += 1
    t7_hdr = ["Família", "Agile cross ΔMAE", "Agile cross ΔPRED25", "Agile LOPO ΔMAE", "Agile LOPO ΔPRED25",
              "NEO cross ΔMAE", "NEO cross ΔPRED25", "NEO LOPO ΔMAE", "NEO LOPO ΔPRED25"]
    for i, h in enumerate(t7_hdr, start=1):
        ws.cell(row=row, column=i, value=h)
    style_range(ws, row, 1, len(t7_hdr), font=WHITE_B, fill=HDR_FILL, align=CENTER)
    t7 = {"ML": [8.1, -11.2, 13.0, -18.0, 13.5, -16.8, 21.0, -26.0],
          "Transformer": [3.4, -3.8, 4.5, -5.0, 5.0, -6.1, 6.2, -6.9],
          "LLM": [0.9, -1.0, 0.8, -1.0, 1.0, -0.9, 1.2, -1.0]}
    for fam, vals in t7.items():
        row += 1
        ws.cell(row=row, column=1, value=fam).font = BOLD
        for i, v in enumerate(vals, start=2):
            c = ws.cell(row=row, column=i, value=v / 100)
            c.font = BLUE
            c.number_format = "+0.0%;-0.0%;0.0%"
        style_range(ws, row, 1, len(t7_hdr))
        ws.cell(row=row, column=1).font = BOLD

    row += 2
    notes = [
        "Leitura: no artigo, todo modelo PIORA ao ir para projetos novos (Δ MAE positivo). Na reprodução, o MAE do ML no "
        "NEO CAI no cross-project. Isso acontece porque os projetos de teste sorteados (Minds Mobile, GitLab Reliability, "
        "Petals Cockpit) usam escalas pequenas (média ~2,6 SP). O resultado depende muito de QUAIS projetos caem no teste, "
        "e o artigo não informa quais usou.",
        "No cross-project, a mediana global (baseline) obtém MAE 1,59 no NEO e 2,71 no Agile, à frente de quase todos os modelos.",
        "Pendente: LOPO do ML (caro: re-treina por projeto), Transformers (Colab) e LLMs.",
        f"Fonte do artigo: {SOURCE_PAPER}, Tabela 7.",
    ]
    for n in notes:
        ws.cell(row=row, column=1, value=n).font = F(italic=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        ws.cell(row=row, column=1).alignment = WRAP
        ws.row_dimensions[row].height = 42
        row += 1
    set_widths(ws, [14, 14, 14, 16, 14, 16, 14, 16, 16])
    ws.row_dimensions[4].height = 30


def repo_sheet(wb):
    ws = wb.create_sheet("Extensao_Repositorios")
    ws["A1"] = "Extensão (não está no artigo): características dos repositórios no NEO"
    ws["A1"].font = TITLE
    ws["A2"] = ("Features: estrelas, forks, issues abertas, tópicos, organização e dias desde a primeira issue do projeto. "
                "MAE em story points; menor é melhor.")
    ws["A2"].font = F(italic=True)
    hdr = ["Modelo", "Entrada", "MAE within", "PRED(25) within", "MAE cross", "PRED(25) cross"]
    r = 4
    for i, h in enumerate(hdr, start=1):
        ws.cell(row=r, column=i, value=h)
    style_range(ws, r, 1, len(hdr), font=WHITE_B, fill=HDR_FILL, align=CENTER)
    rows = [("Mediana do projeto", "baseline (sem texto)", "ProjectMedian"),
            ("Mediana global", "baseline (sem texto)", "Median"),
            ("SVR", "só repositório", "SVR_repoonly"),
            ("XGBoost", "só repositório", "XGBoost_repoonly"),
            ("Random Forest", "só repositório", "RandomForest_repoonly"),
            ("KNN", "só repositório", "KNN_repoonly"),
            ("Linear Regression", "só repositório", "LinearRegression_repoonly"),
            ("XGBoost", "só texto", "XGBoost"),
            ("XGBoost", "texto + repositório", "XGBoost+repo"),
            ("Random Forest", "só texto", "RandomForest"),
            ("Random Forest", "texto + repositório", "RandomForest+repo"),
            ("KNN", "só texto", "KNN"),
            ("KNN", "texto + repositório", "KNN+repo")]
    row = r + 1
    for label, inp, key in rows:
        w = ours("neo", "within", key)
        c = ours("neo", "cross", key)
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=inp)
        vals = [w[0] if w else None, w[3] if w else None, c[0] if c else None, c[3] if c else None]
        for i, v in enumerate(vals, start=3):
            cell = ws.cell(row=row, column=i, value=float(v) if v is not None else "—")
            cell.font = BLUE
            cell.number_format = "0.0%" if i in (4, 6) else "0.000"
            if v is None:
                cell.alignment = CENTER
        style_range(ws, row, 1, len(hdr))
        if inp.startswith("baseline"):
            style_range(ws, row, 1, len(hdr), fill=SUB_FILL)
        row += 1
    row += 1
    notes = [
        "Within: as características do repositório funcionam como identificador do projeto, e o modelo aprende a escala "
        "de story points de cada time. Por isso o SVR só com repositório empata com a mediana do projeto.",
        "Cross: com ~27 projetos de treino, Random Forest e XGBoost decoram os projetos e erram muito em projetos novos "
        "(MAE 17–19). Somadas ao texto, as características não ajudam.",
        "Cuidado: estrelas, forks e issues abertas são um retrato de junho de 2023, posterior à maioria das issues.",
        "Mediana do projeto não existe no cross (o projeto de teste é novo); por isso aparece '—'.",
    ]
    for n in notes:
        ws.cell(row=row, column=1, value=n).font = F(italic=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row=row, column=1).alignment = WRAP
        ws.row_dimensions[row].height = 32
        row += 1
    set_widths(ws, [20, 22, 14, 16, 14, 16])
    ws.row_dimensions[4].height = 30


def lme_sheet(wb):
    ws = wb.create_sheet("LME")
    ws["A1"] = "Modelo linear de efeitos mistos (Seção IV-A.5 do artigo)"
    ws["A1"].font = TITLE
    ws["A2"] = "MAE_ij = μ + β·Família_i + (1 | Modelo_j) + ε_ij, com ML como família de referência."
    ws["A2"].font = F(italic=True)
    hdr = ["Coeficiente", "Artigo", "Nosso script sobre os MAEs das Tabelas 14 e 15", "Diferença",
           "Nosso script sobre os nossos resultados"]
    r = 4
    for i, h in enumerate(hdr, start=1):
        ws.cell(row=r, column=i, value=h)
    style_range(ws, r, 1, len(hdr), font=WHITE_B, fill=HDR_FILL, align=CENTER)
    rows = [("μ (MAE médio do ML)", 5.754, 5.754), ("β Transformer", -1.774, -1.774), ("β LLM", 3.955, 3.955)]
    for k, (lab, a, b) in enumerate(rows, start=r + 1):
        ws.cell(row=k, column=1, value=lab).font = BOLD
        for col, v in [(2, a), (3, b)]:
            c = ws.cell(row=k, column=col, value=v)
            c.font = BLUE
            c.number_format = "0.000"
        ws.cell(row=k, column=4, value=f"=C{k}-B{k}").number_format = "0.000"
        ws.cell(row=k, column=5, value="Pendente (precisa de Transformers e LLMs)")
        style_range(ws, k, 1, len(hdr))
        ws.cell(row=k, column=1).font = BOLD
    k = r + len(rows) + 2
    notes = [
        "O LME do nosso script (91_analysis.py), alimentado com os MAEs publicados nas Tabelas 14 e 15, devolve "
        "exatamente os coeficientes do artigo. Isso confirma que o artigo calculou o LME sobre os valores dessas tabelas.",
        "A variância do efeito aleatório de modelo sai 0: com só 2 observações por modelo (Agile e NEO), o LME "
        "equivale a uma regressão comum com 26 pontos.",
        "O script original do repositório traz um dicionário digitado à mão ('INSERT YOUR RESULTS HERE') com outros "
        "números, em que os LLMs aparecem melhores que os Transformers. Não é a fonte dos coeficientes publicados.",
    ]
    for n in notes:
        ws.cell(row=k, column=1, value=n).font = F(italic=True)
        ws.merge_cells(start_row=k, start_column=1, end_row=k, end_column=5)
        ws.cell(row=k, column=1).alignment = WRAP
        ws.row_dimensions[k].height = 36
        k += 1
    set_widths(ws, [24, 12, 26, 12, 34])
    ws.row_dimensions[4].height = 42


def readme_sheet(wb):
    ws = wb.active
    ws.title = "Leia-me"
    ws["A1"] = "Reprodução de Alaswad et al. (2026) — comparação artigo × nossos resultados"
    ws["A1"].font = TITLE
    lines = [
        ("Artigo", SOURCE_PAPER),
        ("Repositório", "github.com/FeisalAlaswad/SEE-LLM (fork em Story-Points-Estimation, branch reproducao, pasta repro/)"),
        ("", ""),
        ("Abas", ""),
        ("Agile", "Tabela 14 do artigo × reprodução pelo protocolo do texto × código original executado."),
        ("NEO", "Tabela 15 do artigo × reprodução pelo protocolo do texto (o código original não processa o NEO)."),
        ("Generalizacao", "Tabela 7 (cross-project e LOPO) × reprodução."),
        ("Extensao_Repositorios", "Extensão nossa: características dos repositórios do NEO e baselines."),
        ("LME", "Coeficientes do modelo de efeitos mistos: artigo × nosso script."),
        ("", ""),
        ("Legenda", ""),
        ("Texto azul", "Valor copiado: do artigo ou das métricas calculadas pela reprodução."),
        ("Texto preto", "Fórmula (diferenças e médias), recalcula sozinha."),
        ("—", "Não se aplica / não reportado."),
        ("Linhas cinza", "Transformers e LLMs: ainda não rodados na reprodução (Colab e Ollama)."),
        ("", ""),
        ("Principais achados", ""),
        ("1", "Executando o código original, KNN e Regressão Linear reproduzem a Tabela 14 nas três casas decimais; "
              "RF e XGBoost ficam a ≤ 0,02. Os números de ML publicados vêm do código, não do protocolo descrito no texto."),
        ("2", "Pelo protocolo descrito no texto, os erros dos modelos de ML ficam MENORES do que os publicados."),
        ("3", "Prever a mediana de story points do próprio projeto, sem ler o texto, bate quase todos os modelos de ML "
              "e fica muito à frente dos LLMs do artigo. O artigo não compara contra esse baseline."),
        ("4", "O subconjunto 'curado' do NEO é o dataset sem o projeto GitLab: 20.479 stories em 33 projetos (o artigo diz 25)."),
        ("5", "O LME publicado é reproduzido exatamente a partir das Tabelas 14 e 15."),
        ("", ""),
        ("Atualização", "Rode python repro\\90_evaluate.py e gere esta planilha de novo quando houver resultados de "
                        "Transformers e LLMs."),
    ]
    for i, (a, b) in enumerate(lines, start=3):
        ca = ws.cell(row=i, column=1, value=a)
        cb = ws.cell(row=i, column=2, value=b)
        ca.font = BOLD if b == "" or a in ("Artigo", "Repositório", "Atualização") else F(bold=True)
        cb.font = F()
        cb.alignment = WRAP
        if a in ("Abas", "Legenda", "Principais achados"):
            ca.font = Font(name=FONT, size=11, bold=True)
        if a == "Texto azul":
            cb.font = BLUE
        if len(b) > 100:
            ws.row_dimensions[i].height = 14 * (len(b) // 100 + 1)
    set_widths(ws, [22, 110])


wb = Workbook()
readme_sheet(wb)
comparison_sheet(wb, "agile", "Agile", with_code=True)
comparison_sheet(wb, "neo", "NEO", with_code=False)
generalization_sheet(wb)
repo_sheet(wb)
lme_sheet(wb)
wb.save(OUT)
print("ok", OUT)
